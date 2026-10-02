import json
from io import BytesIO
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_POST, require_http_methods
from django.views.decorators.csrf import csrf_exempt
from django.utils import timezone
from users.views import role_required
from quizzes.models import Quiz
from .models import QuizSession, SessionParticipant, LiveAnswer
from .services import (
    active_session_for_teacher,
    get_session_by_code,
    option_stats_for_question,
    pause_session,
    resume_session,
    session_can_accept_student,
    session_participant_results,
    session_leaderboard,
    student_correct_answers,
    teacher_can_access_quiz,
)


# ─── Helper ──────────────────────────────────────────────────────────────────

def _auto_advance_if_needed(session):
    """Auto-advance question if timer expired. Returns True if advanced."""
    if session.estado != QuizSession.EN_CURSO:
        return False
    if not session.ultima_pregunta_inicio:
        return False

    tiempo_restante = session.tiempo_restante()
    if tiempo_restante > 0:
        return False

    # Timer expired → advance
    total = session.total_preguntas()
    if session.pregunta_actual >= total:
        # End session
        session.estado = QuizSession.FINALIZADO
        session.fecha_fin = timezone.now()
        session.save()
    else:
        session.pregunta_actual += 1
        session.ultima_pregunta_inicio = timezone.now()
        session.save()
    return True


# ─── Docente: Session Management ─────────────────────────────────────────────

@role_required('DOCENTE')
def create_session(request):
    # Rule: One active session per teacher
    active_session = active_session_for_teacher(request.user)
    
    if active_session:
        messages.warning(request, 'Ya tienes una sesión activa. Termínala primero para empezar un nuevo Quiz.')
        if active_session.estado == QuizSession.EN_CURSO:
            return redirect('live_session_docente', session_id=active_session.pk)
        return redirect('waiting_room_docente', session_id=active_session.pk)

    if request.method == 'POST':
        quiz_id = request.POST.get('quiz_id')
        quiz = get_object_or_404(Quiz, pk=quiz_id, publicado=True)
        if not teacher_can_access_quiz(request.user, quiz):
            messages.error(request, 'No tienes acceso a ese quiz.')
            return redirect('quiz_list')

        session = QuizSession.objects.create(quiz=quiz, docente=request.user)
        messages.success(request, f'Sesión creada con código: {session.codigo}')
        return redirect('waiting_room_docente', session_id=session.pk)
    
    return redirect('quiz_list')


@role_required('DOCENTE')
def session_list(request):
    sessions = QuizSession.objects.filter(docente=request.user).select_related('quiz').order_by('-id')
    return render(request, 'sessions_app/session_list.html', {'sessions': sessions})


@role_required('DOCENTE')
def waiting_room_docente(request, session_id):
    session = get_object_or_404(QuizSession, pk=session_id, docente=request.user)
    return render(request, 'sessions_app/waiting_room_docente.html', {'session': session})


@role_required('DOCENTE')
def live_session_docente(request, session_id):
    session = get_object_or_404(QuizSession, pk=session_id, docente=request.user)
    if session.estado == QuizSession.ESPERA:
        return redirect('waiting_room_docente', session_id=session_id)
    return render(request, 'sessions_app/live_session_docente.html', {'session': session})


# ─── Estudiante: Join/Play ────────────────────────────────────────────────────

@role_required('ESTUDIANTE')
def join_session(request):
    if request.method == 'POST':
        codigo = request.POST.get('codigo', '').strip().upper()
        session = get_session_by_code(codigo)
        if not session:
            messages.error(request, 'Código de sesión no encontrado.')
            return redirect('estudiante_dashboard')

        can_join, join_message = session_can_accept_student(session, request.user)
        if not can_join:
            messages.error(request, join_message)
            return redirect('estudiante_dashboard')

        participant, created = SessionParticipant.objects.get_or_create(
            session=session, student=request.user,
            defaults={'activo': True, 'expulsado': False}
        )
        if participant.expulsado:
            messages.error(request, 'Fuiste expulsado de esta sesión.')
            return redirect('estudiante_dashboard')

        participant.activo = True
        participant.save()

        if session.estado in (QuizSession.EN_CURSO, QuizSession.PAUSADO):
            return redirect('live_session_estudiante', session_id=session.pk)
        return redirect('waiting_room_estudiante', session_id=session.pk)

    return render(request, 'sessions_app/join_session.html')


@role_required('ESTUDIANTE')
def waiting_room_estudiante(request, session_id):
    session = get_object_or_404(QuizSession, pk=session_id)
    participant = get_object_or_404(SessionParticipant, session=session, student=request.user)
    if participant.expulsado:
        messages.error(request, 'Fuiste expulsado de esta sesión.')
        return redirect('estudiante_dashboard')
    return render(request, 'sessions_app/waiting_room_estudiante.html', {
        'session': session, 'participant': participant
    })


@role_required('ESTUDIANTE')
def live_session_estudiante(request, session_id):
    session = get_object_or_404(QuizSession, pk=session_id)
    try:
        participant = SessionParticipant.objects.get(session=session, student=request.user)
    except SessionParticipant.DoesNotExist:
        return redirect('join_session')

    if participant.expulsado:
        messages.error(request, 'Fuiste expulsado de esta sesión.')
        return redirect('estudiante_dashboard')

    return render(request, 'sessions_app/live_session_estudiante.html', {
        'session': session,
        'participant': participant
    })


# ─── AJAX Polling API ─────────────────────────────────────────────────────────

def session_state_api(request, session_id):
    """Polling endpoint — returns full session state."""
    session = get_object_or_404(QuizSession, pk=session_id)

    # Auto-advance if timer expired
    _auto_advance_if_needed(session)
    session.refresh_from_db()

    pregunta_obj = session.pregunta_actual_obj()
    participants = session.participantes_activos()
    participant_count = participants.count()

    # Count answers for current question
    answers_count = 0
    if pregunta_obj:
        answers_count = LiveAnswer.objects.filter(
            session=session, question=pregunta_obj
        ).count()

    # For enrolled student: did they answer?
    student_answered = False
    student_expulsado = False
    if request.user.is_authenticated and hasattr(request.user, 'profile'):
        if request.user.profile.role == 'ESTUDIANTE':
            try:
                p = SessionParticipant.objects.get(session=session, student=request.user)
                student_expulsado = p.expulsado
                if pregunta_obj:
                    student_answered = LiveAnswer.objects.filter(
                        session=session, student=request.user, question=pregunta_obj
                    ).exists()
            except SessionParticipant.DoesNotExist:
                pass

    data = {
        'estado': session.estado,
        'pregunta_actual': session.pregunta_actual,
        'total_preguntas': session.total_preguntas(),
        'quiz_tiempo': session.quiz.tiempo_por_pregunta,
        'tiempo_restante': session.tiempo_restante(),
        'participant_count': participant_count,
        'answers_count': answers_count,
        'permitir_ingreso': session.permitir_ingreso,
        'student_answered': student_answered,
        'student_expulsado': student_expulsado,
        'mostrando_resultados': session.mostrando_resultados,
    }

    if session.mostrando_resultados and student_answered and pregunta_obj:
        user_ans = LiveAnswer.objects.filter(session=session, student=request.user, question=pregunta_obj).select_related('option').first()
        if user_ans:
            correct_opts = [o.letra for o in pregunta_obj.options.filter(es_correcta=True)]
            data['student_result'] = {
                'es_correcta': user_ans.option.es_correcta,
                'puntaje': user_ans.puntaje,
                'letra_marcada': user_ans.option.letra,
                'correct_opts': correct_opts
            }

    if pregunta_obj:
        data['pregunta'] = {
            'id': pregunta_obj.pk,
            'enunciado': pregunta_obj.enunciado,
            'orden': pregunta_obj.orden,
            'imagen': pregunta_obj.imagen.url if pregunta_obj.imagen else None,
            'opciones': [
                {'id': o.pk, 'letra': o.letra, 'texto': o.texto}
                for o in pregunta_obj.options.order_by('letra')
            ]
        }
    else:
        data['pregunta'] = None

    return JsonResponse(data)


def participants_api(request, session_id):
    """Returns list of active participants."""
    session = get_object_or_404(QuizSession, pk=session_id)
    participants = session.participantes_activos().select_related('student')

    pregunta_obj = session.pregunta_actual_obj()
    answered_ids = set()
    if pregunta_obj:
        answered_ids = set(LiveAnswer.objects.filter(
            session=session, question=pregunta_obj
        ).values_list('student_id', flat=True))

    data = {
        'participants': [
            {
                'id': p.student.pk,
                'username': p.student.username,
                'answered': p.student.pk in answered_ids,
            }
            for p in participants
        ]
    }
    return JsonResponse(data)


def results_api(request, session_id, question_orden):
    """Returns vote counts per option for a given question (by orden)."""
    session = get_object_or_404(QuizSession, pk=session_id)
    questions = list(session.quiz.questions.prefetch_related('options').order_by('orden'))

    idx = int(question_orden) - 1
    if idx < 0 or idx >= len(questions):
        return JsonResponse({'error': 'Pregunta no encontrada'}, status=404)

    question = questions[idx]
    result, total_answers = option_stats_for_question(session, question)
    return JsonResponse({'opciones': result, 'total': total_answers})


# ─── Docente Actions ──────────────────────────────────────────────────────────

@require_POST
@role_required('DOCENTE')
def start_session(request, session_id):
    session = get_object_or_404(QuizSession, pk=session_id, docente=request.user)
    if session.estado != QuizSession.ESPERA:
        return JsonResponse({'error': 'La sesión ya fue iniciada.'}, status=400)

    session.estado = QuizSession.EN_CURSO
    session.fecha_inicio = timezone.now()
    session.pregunta_actual = 1
    session.ultima_pregunta_inicio = timezone.now()
    session.save()
    return JsonResponse({'ok': True, 'session_id': session.pk})


@require_POST
@role_required('DOCENTE')
def next_question(request, session_id):
    session = get_object_or_404(QuizSession, pk=session_id, docente=request.user)
    if session.estado != QuizSession.EN_CURSO:
        return JsonResponse({'error': 'La sesión no está en curso.'}, status=400)

    force = request.POST.get('force', 'false').lower() == 'true'
    pregunta_obj = session.pregunta_actual_obj()

    if pregunta_obj:
        participants = session.participantes_activos()
        total_active = participants.count()
        answered = LiveAnswer.objects.filter(session=session, question=pregunta_obj).count()
        pending = total_active - answered

        if pending > 0 and not force:
            return JsonResponse({'needs_confirm': True, 'pending': pending})

    total = session.total_preguntas()
    if session.pregunta_actual >= total:
        session.estado = QuizSession.FINALIZADO
        session.fecha_fin = timezone.now()
        session.save()
        return JsonResponse({'ok': True, 'finalizado': True})
    else:
        session.pregunta_actual += 1
        session.ultima_pregunta_inicio = timezone.now()
        session.save()
        return JsonResponse({'ok': True, 'finalizado': False, 'pregunta_actual': session.pregunta_actual})


@require_POST
@role_required('DOCENTE')
def close_admission(request, session_id):
    session = get_object_or_404(QuizSession, pk=session_id, docente=request.user)
    session.permitir_ingreso = False
    session.save()
    return JsonResponse({'ok': True})


@require_POST
@role_required('DOCENTE')
def expel_student(request, session_id, student_id):
    session = get_object_or_404(QuizSession, pk=session_id, docente=request.user)
    participant = get_object_or_404(SessionParticipant, session=session, student_id=student_id)
    participant.expulsado = True
    participant.activo = False
    participant.save()
    return JsonResponse({'ok': True})


@require_POST
@role_required('DOCENTE')
def end_session(request, session_id):
    session = get_object_or_404(QuizSession, pk=session_id, docente=request.user)
    session.estado = QuizSession.FINALIZADO
    session.fecha_fin = timezone.now()
    session.permitir_ingreso = False
    session.save()
    return JsonResponse({'ok': True})


@require_POST
@role_required('DOCENTE')
def pause_session_view(request, session_id):
    session = get_object_or_404(QuizSession, pk=session_id, docente=request.user)
    if not pause_session(session):
        return JsonResponse({'error': 'La sesión no se puede pausar.'}, status=400)
    return JsonResponse({'ok': True, 'estado': session.estado})


@require_POST
@role_required('DOCENTE')
def resume_session_view(request, session_id):
    session = get_object_or_404(QuizSession, pk=session_id, docente=request.user)
    if not resume_session(session):
        return JsonResponse({'error': 'La sesión no se puede reanudar.'}, status=400)
    return JsonResponse({'ok': True, 'estado': session.estado})


# ─── Estudiante Actions ───────────────────────────────────────────────────────

@require_POST
@role_required('ESTUDIANTE')
def submit_answer(request, session_id):
    session = get_object_or_404(QuizSession, pk=session_id)

    if session.estado == QuizSession.PAUSADO:
        return JsonResponse({'error': 'La sesión está pausada.'}, status=400)

    if session.estado != QuizSession.EN_CURSO:
        return JsonResponse({'error': 'Sesión no activa.'}, status=400)

    try:
        participant = SessionParticipant.objects.get(session=session, student=request.user)
    except SessionParticipant.DoesNotExist:
        return JsonResponse({'error': 'No participas en esta sesión.'}, status=403)

    if participant.expulsado:
        return JsonResponse({'error': 'Fuiste expulsado.'}, status=403)

    # Check timer
    if session.tiempo_restante() == 0:
        return JsonResponse({'error': 'El tiempo ha expirado.'}, status=400)

    pregunta_obj = session.pregunta_actual_obj()
    if not pregunta_obj:
        return JsonResponse({'error': 'No hay pregunta activa.'}, status=400)

    # Check not already answered
    if LiveAnswer.objects.filter(session=session, student=request.user, question=pregunta_obj).exists():
        return JsonResponse({'error': 'Ya respondiste esta pregunta.'}, status=400)

    option_id = request.POST.get('option_id')
    from quizzes.models import Option
    option = get_object_or_404(Option, pk=option_id, question=pregunta_obj)

    puntaje = 0
    if option.es_correcta:
        elapsed = (timezone.now() - session.ultima_pregunta_inicio).total_seconds()
        tiempo_total = session.quiz.tiempo_por_pregunta
        respuestas_previas = LiveAnswer.objects.filter(
            session=session, question=pregunta_obj, option__es_correcta=True
        ).count()
        
        if respuestas_previas == 0:
            puntaje = 1000
        else:
            penalizacion_tiempo = int((elapsed / tiempo_total) * 500) if tiempo_total > 0 else 0
            penalizacion_posicion = 50 * respuestas_previas
            puntaje = max(100, 1000 - penalizacion_posicion - penalizacion_tiempo)

    LiveAnswer.objects.create(
        session=session, student=request.user,
        question=pregunta_obj, option=option, puntaje=puntaje
    )
    return JsonResponse({'ok': True})


def leaderboard_api(request, session_id):
    """Devuelve el leaderboard parcial de toda la sesión (top 6 puntajes altos)."""
    session = get_object_or_404(QuizSession, pk=session_id)
    return JsonResponse({'ranking': session_leaderboard(session)})


# ─── Session Results ──────────────────────────────────────────────────────────

def session_results(request, session_id):
    session = get_object_or_404(QuizSession, pk=session_id)
    questions = list(session.quiz.questions.prefetch_related('options').order_by('orden'))
    hidden_columns = [item for item in request.GET.get('hide', '').split(',') if item]
    participant_results = session_participant_results(session)

    results = []
    for q in questions:
        opts, _total_votos_q = option_stats_for_question(session, q)
        results.append({'pregunta': q, 'opciones': opts})

    # Student score
    student_score = None
    if request.user.is_authenticated and hasattr(request.user, 'profile'):
        if request.user.profile.role == 'ESTUDIANTE':
            correct = student_correct_answers(session, request.user)
            student_score = {'correctas': correct, 'total': len(questions)}

    return render(request, 'sessions_app/session_results.html', {
        'session': session,
        'results': results,
        'student_score': student_score,
        'leaderboard': session_leaderboard(session, limit=5),
        'participant_results': participant_results,
        'hidden_columns': hidden_columns,
    })


@require_http_methods(['GET'])
@role_required('DOCENTE')
def export_results_xlsx(request, session_id):
    session = get_object_or_404(QuizSession, pk=session_id, docente=request.user)
    hidden_columns = [item for item in request.GET.get('hide', '').split(',') if item]
    participant_results = session_participant_results(session)
    questions = list(session.quiz.questions.prefetch_related('options').order_by('orden'))

    try:
        from openpyxl import Workbook
    except ImportError:
        messages.error(request, 'Falta instalar openpyxl para exportar a XLSX.')
        return redirect('session_results', session_id=session_id)

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = 'Resumen'
    summary_sheet.append(['Quiz', session.quiz.titulo])
    summary_sheet.append(['Sesión', session.codigo])
    summary_sheet.append(['Participantes', len(participant_results)])
    summary_sheet.append(['Preguntas', len(questions)])

    participant_sheet = workbook.create_sheet('Participantes')
    columns = [
        ('username', 'Estudiante'),
        ('score', 'Puntaje'),
        ('correctas', 'Correctas'),
        ('respondidas', 'Respondidas'),
        ('porcentaje', '% Acierto'),
    ]
    visible_columns = [column for column in columns if column[0] not in hidden_columns]
    participant_sheet.append([label for _, label in visible_columns])
    for row in participant_results:
        participant_sheet.append([row[key] for key, _ in visible_columns])

    question_sheet = workbook.create_sheet('Preguntas')
    question_sheet.append(['Orden', 'Pregunta', 'Correcta', 'Explicación'])
    for question in questions:
        correct = question.correct_option()
        question_sheet.append([
            question.orden,
            question.enunciado,
            correct.letra if correct else '',
            question.explicacion or '',
        ])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="quiz-{session.codigo}-resultados.xlsx"'
    return response


@require_http_methods(['GET'])
@role_required('DOCENTE')
def export_results_pdf(request, session_id):
    session = get_object_or_404(QuizSession, pk=session_id, docente=request.user)
    hidden_columns = [item for item in request.GET.get('hide', '').split(',') if item]
    participant_results = session_participant_results(session)
    questions = list(session.quiz.questions.prefetch_related('options').order_by('orden'))

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    except ImportError:
        messages.error(request, 'Falta instalar reportlab para exportar a PDF.')
        return redirect('session_results', session_id=session_id)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=1.5 * cm, leftMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(f'Resultados del quiz: {session.quiz.titulo}', styles['Title']))
    story.append(Paragraph(f'Sesión: {session.codigo}', styles['Normal']))
    story.append(Paragraph(f'Participantes: {len(participant_results)} | Preguntas: {len(questions)}', styles['Normal']))
    story.append(Spacer(1, 12))

    participant_columns = [
        ('username', 'Estudiante'),
        ('score', 'Puntaje'),
        ('correctas', 'Correctas'),
        ('respondidas', 'Respondidas'),
        ('porcentaje', '% Acierto'),
    ]
    visible_columns = [column for column in participant_columns if column[0] not in hidden_columns]
    if visible_columns:
        table_data = [[label for _, label in visible_columns]]
        for row in participant_results:
            table_data.append([str(row[key]) for key, _ in visible_columns])
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f3c88')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ]))
        story.append(Paragraph('Participantes', styles['Heading2']))
        story.append(table)
        story.append(Spacer(1, 12))

    story.append(Paragraph('Preguntas', styles['Heading2']))
    for question in questions:
        correct = question.correct_option()
        table = Table([
            ['Orden', 'Pregunta', 'Correcta', 'Explicación'],
            [str(question.orden), question.enunciado, correct.letra if correct else '', question.explicacion or ''],
        ], repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4b5563')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ]))
        story.append(table)
        story.append(Spacer(1, 10))

    doc.build(story)
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="quiz-{session.codigo}-resultados.pdf"'
    response.write(buffer.getvalue())
    buffer.close()
    return response
