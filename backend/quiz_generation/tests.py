from io import BytesIO
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.test import override_settings
from django.utils import timezone
from datetime import timedelta
from django.urls import reverse

from courses.models import Course, CourseTeacher
from quizzes.models import Quiz
from users.models import Profile

from .extraction import extract_text
from .generator import generate_quiz
from .models import GenerationJob, SourceDocument
from .providers import ProviderError, ProviderNotConfigured, generate_with_retries, get_provider


class FakeProvider:
    def __init__(self, payload):
        self.payload = payload

    def generate(self, prompt, schema):
        return self.payload


def question(number, statement=None, kind='MC'):
    return {
        'tipo': kind,
        'enunciado': statement or f'Pregunta {number}',
        'opciones': {'A': 'Correcta', 'B': 'Incorrecta', 'C': 'Otra', 'D': 'Otra más'} if kind == 'MC' else {'A': 'Verdadero', 'B': 'Falso'},
        'correcta': 'A',
        'explicacion': 'Explicación',
    }


class ExtractionTests(TestCase):
    def test_extract_text_txt_and_reject_short_source(self):
        source = BytesIO(('Contenido suficiente ' * 20).encode())
        self.assertIn('Contenido suficiente', extract_text(source, 'TXT'))
        with self.assertRaisesMessage(ValueError, 'No se pudo extraer texto'):
            extract_text(BytesIO(b'corto'), 'TXT')

    def test_extract_text_from_docx_pptx_and_pdf(self):
        from docx import Document
        from pptx import Presentation
        from reportlab.pdfgen import canvas

        docx_stream = BytesIO()
        document = Document()
        document.add_paragraph('DOCX ' * 50)
        document.save(docx_stream)
        self.assertIn('DOCX', extract_text(docx_stream, 'DOCX'))

        pptx_stream = BytesIO()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = 'PPTX ' * 50
        presentation.save(pptx_stream)
        self.assertIn('PPTX', extract_text(pptx_stream, 'PPTX'))

        pdf_stream = BytesIO()
        pdf = canvas.Canvas(pdf_stream)
        pdf.drawString(10, 800, 'PDF ' * 60)
        pdf.save()
        self.assertIn('PDF', extract_text(pdf_stream, 'PDF'))


class ProviderTests(TestCase):
    def test_retries_then_returns_provider_result(self):
        class Flaky:
            calls = 0

            def generate(self, prompt, schema):
                self.calls += 1
                if self.calls == 1:
                    raise RuntimeError('temporary')
                return {'ok': True}

        with patch('quiz_generation.providers.time.sleep'):
            self.assertEqual(generate_with_retries(Flaky(), 'prompt', {}), {'ok': True})

    def test_quota_error_has_user_message(self):
        class Quota:
            def generate(self, prompt, schema):
                raise RuntimeError('429 quota exceeded')

        with patch('quiz_generation.providers.time.sleep'):
            with self.assertRaisesMessage(ProviderError, 'cuota gratuita'):
                generate_with_retries(Quota(), 'prompt', {})

    @override_settings(QUIZ_AI_PROVIDER='', GEMINI_API_KEY='', DEEPSEEK_API_KEY='')
    def test_missing_provider_is_explicit(self):
        with self.assertRaisesMessage(ProviderNotConfigured, 'no está configurada'):
            get_provider()

    @override_settings(QUIZ_AI_PROVIDER='gemini', GEMINI_API_KEY='key', GEMINI_MODEL='model')
    def test_gemini_provider_is_selected_from_settings(self):
        with patch('quiz_generation.providers.GeminiProvider') as provider:
            self.assertIs(provider.return_value, get_provider())
            provider.assert_called_once_with('key', 'model')

    @override_settings(QUIZ_AI_PROVIDER='deepseek', DEEPSEEK_API_KEY='key', DEEPSEEK_MODEL='model')
    def test_deepseek_provider_is_selected_from_settings(self):
        with patch('quiz_generation.providers.DeepSeekProvider') as provider:
            self.assertIs(provider.return_value, get_provider())
            provider.assert_called_once_with('key', 'model')


class GeneratorTests(TestCase):
    def setUp(self):
        self.teacher = User.objects.create_user(username='generation_teacher', password='password123')
        Profile.objects.filter(user=self.teacher).update(role=Profile.DOCENTE)
        self.course = Course.objects.create(nombre='Generación', descripcion='', creado_por=self.teacher)
        CourseTeacher.objects.create(course=self.course, teacher=self.teacher)
        document = SourceDocument.objects.create(
            course=self.course,
            docente=self.teacher,
            archivo=SimpleUploadedFile('fuente.txt', ('Texto fuente ' * 30).encode()),
            nombre_original='fuente.txt',
            tipo=SourceDocument.TXT,
            texto_extraido='Texto fuente ' * 30,
        )
        self.job = GenerationJob.objects.create(
            docente=self.teacher,
            course=self.course,
            parametros={'titulo': 'Quiz generado', 'tipo': 'MIXTO', 'instrucciones': ''},
            preguntas_solicitadas=10,
        )
        self.job.documentos.add(document)

    def test_valid_questions_create_unpublished_quiz(self):
        payload = {'preguntas': [question(index) for index in range(1, 6)] + [question(6, 'Pregunta 1')]}
        with patch('quiz_generation.generator.get_provider', return_value=FakeProvider(payload)):
            generate_quiz(self.job.pk)
        self.job.refresh_from_db()
        self.assertEqual(self.job.estado, GenerationJob.COMPLETADO)
        self.assertEqual(self.job.preguntas_generadas, 5)
        self.assertFalse(Quiz.objects.get(pk=self.job.quiz_id).publicado)

    def test_less_than_five_valid_questions_fails(self):
        payload = {'preguntas': [question(index) for index in range(1, 4)]}
        with patch('quiz_generation.generator.get_provider', return_value=FakeProvider(payload)):
            generate_quiz(self.job.pk)
        self.job.refresh_from_db()
        self.assertEqual(self.job.estado, GenerationJob.FALLIDO)
        self.assertIn('menos de 5', self.job.mensaje)

    def test_missing_provider_configuration_fails_without_500(self):
        with patch('quiz_generation.generator.get_provider', side_effect=ProviderNotConfigured('La generación con IA no está configurada.')):
            generate_quiz(self.job.pk)
        self.job.refresh_from_db()
        self.assertEqual(self.job.estado, GenerationJob.FALLIDO)
        self.assertIn('no está configurada', self.job.mensaje)


class GenerationViewTests(TestCase):
    def setUp(self):
        self.teacher = User.objects.create_user(username='view_generation_teacher', password='password123')
        self.other = User.objects.create_user(username='other_generation_teacher', password='password123')
        Profile.objects.filter(user=self.teacher).update(role=Profile.DOCENTE)
        Profile.objects.filter(user=self.other).update(role=Profile.DOCENTE)
        self.course = Course.objects.create(nombre='Curso propio', descripcion='', creado_por=self.teacher)
        self.other_course = Course.objects.create(nombre='Curso ajeno', descripcion='', creado_por=self.other)
        CourseTeacher.objects.create(course=self.course, teacher=self.teacher)
        CourseTeacher.objects.create(course=self.other_course, teacher=self.other)

    def test_other_course_is_not_available(self):
        self.client.force_login(self.teacher)
        response = self.client.post(reverse('generate_quiz'), {
            'titulo': 'Ajeno', 'course': self.other_course.pk,
            'preguntas_solicitadas': 5, 'tipo': 'MIXTO', 'instrucciones': '',
        })
        self.assertEqual(response.status_code, 200)
        self.assertFalse(GenerationJob.objects.exists())

    def test_generation_view_starts_job_after_commit(self):
        source = SimpleUploadedFile('fuente.txt', ('Contenido ' * 60).encode(), content_type='text/plain')
        self.client.force_login(self.teacher)
        with patch('quiz_generation.views._start_job') as start_job:
            with self.captureOnCommitCallbacks(execute=True):
                response = self.client.post(reverse('generate_quiz'), {
                    'titulo': 'Desde archivo', 'course': self.course.pk,
                    'preguntas_solicitadas': 5, 'tipo': 'MIXTO', 'instrucciones': '',
                    'archivos': source,
                })
        self.assertEqual(response.status_code, 302)
        self.assertTrue(GenerationJob.objects.exists())
        start_job.assert_called_once()

    def test_invalid_file_is_rejected(self):
        source = SimpleUploadedFile('fuente.exe', b'contenido', content_type='application/octet-stream')
        self.client.force_login(self.teacher)
        response = self.client.post(reverse('generate_quiz'), {
            'titulo': 'Invalido', 'course': self.course.pk,
            'preguntas_solicitadas': 5, 'tipo': 'MIXTO', 'instrucciones': '',
            'archivos': source,
        })
        self.assertEqual(response.status_code, 200)
        self.assertFalse(GenerationJob.objects.exists())

    def test_status_is_private_and_marks_stale_job_failed(self):
        job = GenerationJob.objects.create(
            docente=self.teacher,
            course=self.course,
            parametros={'titulo': 'Stale'},
            estado=GenerationJob.PROCESANDO,
            iniciado_en=timezone.now() - timedelta(minutes=11),
        )
        self.client.force_login(self.teacher)
        response = self.client.get(reverse('generation_status_json', args=[job.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['estado'], GenerationJob.FALLIDO)
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(reverse('generation_status', args=[job.pk])).status_code, 404)

    def test_import_json_valid_text_creates_quiz_and_redirects(self):
        self.client.force_login(self.teacher)
        json_payload = '''{
            "preguntas": [
                {
                    "tipo": "MC",
                    "enunciado": "¿Cuál es la capital de Colombia?",
                    "opciones": {"A": "Bogotá", "B": "Cali", "C": "Medellín", "D": "Barranquilla"},
                    "correcta": "A",
                    "explicacion": "Bogotá es la capital."
                },
                {
                    "tipo": "VF",
                    "enunciado": "La tierra es plana.",
                    "opciones": {"A": "Verdadero", "B": "Falso"},
                    "correcta": "B",
                    "explicacion": "Es un esferoide oblato."
                }
            ]
        }'''
        response = self.client.post(reverse('import_quiz_json'), {
            'titulo': 'Quiz Importado',
            'course': self.course.pk,
            'descripcion': 'Probando importación JSON',
            'json_text': json_payload,
        })
        self.assertEqual(response.status_code, 302)
        quiz = Quiz.objects.get(titulo='Quiz Importado')
        self.assertEqual(quiz.questions.count(), 2)
        self.assertFalse(quiz.publicado)

    def test_import_json_handles_markdown_codeblock_and_file_upload(self):
        self.client.force_login(self.teacher)
        json_payload = '''```json
        {
            "preguntas": [
                {
                    "tipo": "MC",
                    "enunciado": "Pregunta con bloque markdown",
                    "opciones": {"A": "A", "B": "B", "C": "C", "D": "D"},
                    "correcta": "A",
                    "explicacion": "Explicación"
                }
            ]
        }
        ```'''
        json_file = SimpleUploadedFile('quiz.json', json_payload.encode('utf-8'), content_type='application/json')
        response = self.client.post(reverse('import_quiz_json'), {
            'titulo': 'Quiz desde Archivo JSON',
            'course': self.course.pk,
            'archivo_json': json_file,
        })
        self.assertEqual(response.status_code, 302)
        quiz = Quiz.objects.get(titulo='Quiz desde Archivo JSON')
        self.assertEqual(quiz.questions.count(), 1)

    def test_import_json_invalid_structure_shows_form_error(self):
        self.client.force_login(self.teacher)
        response = self.client.post(reverse('import_quiz_json'), {
            'titulo': 'Quiz Inválido',
            'course': self.course.pk,
            'json_text': '{"preguntas": "no es una lista"}',
        })
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Quiz.objects.filter(titulo='Quiz Inválido').exists())

