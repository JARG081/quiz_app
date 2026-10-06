from pathlib import Path


def extract_text(file_object, document_type):
    file_object.seek(0)
    if document_type == 'PDF':
        from pypdf import PdfReader
        text = '\n'.join(page.extract_text() or '' for page in PdfReader(file_object).pages)
    elif document_type == 'DOCX':
        from docx import Document
        text = '\n'.join(paragraph.text for paragraph in Document(file_object).paragraphs)
    elif document_type == 'PPTX':
        from pptx import Presentation
        text = '\n'.join(
            shape.text
            for slide in Presentation(file_object).slides
            for shape in slide.shapes
            if hasattr(shape, 'text')
        )
    elif document_type == 'TXT':
        text = file_object.read().decode('utf-8', errors='replace')
    else:
        raise ValueError('Tipo de archivo no soportado.')
    text = text.strip()
    if len(text) < 200:
        raise ValueError('No se pudo extraer texto (¿PDF escaneado?)')
    return text


def document_type_from_name(name):
    extension = Path(name).suffix.lower()
    types = {'.pdf': 'PDF', '.docx': 'DOCX', '.pptx': 'PPTX', '.txt': 'TXT'}
    if extension not in types:
        raise ValueError('Formato no permitido. Usa PDF, DOCX, PPTX o TXT.')
    return types[extension]
